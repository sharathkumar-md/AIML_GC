"""
Chart builder for creating native PowerPoint charts.
"""

import logging
from typing import Dict, List, Any, Optional

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor as RgbColor  # RGBColor is in pptx.dml.color

logger = logging.getLogger("chart_builder")


class ChartBuilder:
    """
    Builds native PowerPoint charts from financial data.
    """

    def __init__(self, branding=None):
        """
        Initialize chart builder.

        Args:
            branding: Kelp branding configuration
        """
        self.branding = branding

        # Chart colors
        self.colors = [
            RgbColor(45, 35, 90),      # Primary dark
            RgbColor(0, 200, 255),      # Cyan
            RgbColor(255, 100, 130),    # Pink
            RgbColor(255, 150, 80),     # Orange
        ]

    def create_bar_chart(
        self,
        slide,
        data: Dict[str, List[Dict]],
        position: tuple,
        size: tuple,
        title: str = ""
    ):
        """
        Create a bar chart on the slide.

        Args:
            slide: PowerPoint slide
            data: Dictionary with series name as key, list of {year, value} as value
            position: (left, top) in inches
            size: (width, height) in inches
            title: Chart title
        """
        logger.info(f"Creating bar chart: {title}")

        chart_data = CategoryChartData()

        # Get categories (years) from first series
        first_series = list(data.values())[0] if data else []
        categories = [d.get('year', '') for d in first_series]
        chart_data.categories = categories

        # Add each series
        for series_name, series_data in data.items():
            values = [d.get('value', 0) for d in series_data]
            chart_data.add_series(series_name, values)

        # Add chart to slide
        left, top = position
        width, height = size

        chart_shape = slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            Inches(left), Inches(top),
            Inches(width), Inches(height),
            chart_data
        )

        chart = chart_shape.chart

        # Style the chart
        self._style_chart(chart, title)

        return chart

    def create_line_chart(
        self,
        slide,
        data: Dict[str, List[Dict]],
        position: tuple,
        size: tuple,
        title: str = ""
    ):
        """
        Create a line chart on the slide.

        Args:
            slide: PowerPoint slide
            data: Dictionary with series data
            position: (left, top) in inches
            size: (width, height) in inches
            title: Chart title
        """
        logger.info(f"Creating line chart: {title}")

        chart_data = CategoryChartData()

        first_series = list(data.values())[0] if data else []
        categories = [d.get('year', '') for d in first_series]
        chart_data.categories = categories

        for series_name, series_data in data.items():
            values = [d.get('value', 0) for d in series_data]
            chart_data.add_series(series_name, values)

        left, top = position
        width, height = size

        chart_shape = slide.shapes.add_chart(
            XL_CHART_TYPE.LINE_MARKERS,
            Inches(left), Inches(top),
            Inches(width), Inches(height),
            chart_data
        )

        chart = chart_shape.chart
        self._style_chart(chart, title)

        return chart

    def create_pie_chart(
        self,
        slide,
        data: List[Dict],
        position: tuple,
        size: tuple,
        title: str = ""
    ):
        """
        Create a pie chart on the slide.

        Args:
            slide: PowerPoint slide
            data: List of {label, value} dictionaries
            position: (left, top) in inches
            size: (width, height) in inches
            title: Chart title
        """
        logger.info(f"Creating pie chart: {title}")

        chart_data = CategoryChartData()
        categories = [d.get('label', '') for d in data]
        values = [d.get('value', 0) for d in data]

        chart_data.categories = categories
        chart_data.add_series('Values', values)

        left, top = position
        width, height = size

        chart_shape = slide.shapes.add_chart(
            XL_CHART_TYPE.PIE,
            Inches(left), Inches(top),
            Inches(width), Inches(height),
            chart_data
        )

        chart = chart_shape.chart
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.RIGHT

        return chart

    def _style_chart(self, chart, title: str = "") -> None:
        """Apply styling to a chart."""
        # Legend
        chart.has_legend = True
        chart.legend.include_in_layout = False
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM

        # Title
        if title:
            chart.has_title = True
            chart.chart_title.text_frame.paragraphs[0].text = title
            chart.chart_title.text_frame.paragraphs[0].font.size = Pt(12)

    def prepare_financial_chart_data(
        self,
        financials,
        metrics: List[str] = ['revenue', 'ebitda']
    ) -> Dict[str, List[Dict]]:
        """
        Prepare financial data for charting.

        Args:
            financials: Financials object
            metrics: List of metrics to include

        Returns:
            Dictionary ready for chart creation
        """
        data = {}

        # Get sorted years
        years = sorted(financials.yearly_data.keys())[-6:]  # Last 6 years

        for metric in metrics:
            series_data = []
            for year in years:
                fy = financials.yearly_data.get(year)
                if fy:
                    value = getattr(fy, metric, None)
                    if value is not None:
                        series_data.append({
                            'year': f"FY{str(year)[2:]}",
                            'value': round(value, 1)
                        })

            if series_data:
                # Capitalize metric name for display
                display_name = metric.upper() if metric in ['ebitda', 'pat'] else metric.title()
                data[display_name] = series_data

        return data


def create_chart_builder(branding=None) -> ChartBuilder:
    """Factory function to create a ChartBuilder."""
    return ChartBuilder(branding)
