"""
Kelp Automated Deal Flow - Main Entry Point
AI-ML GC 2025-26

This is the main entry point for generating Investment Teaser presentations.
"""

import argparse
import sys
from pathlib import Path
from typing import Optional, List

# Add src to path
sys.path.insert(0, str(Path(__file__).parent))

from utils.logger import setup_logger, get_logger, LogTimer
from config import get_config, get_sector_template, PROJECT_CODENAMES
from ingestion.data_loader import DataLoader
from models.company_data import CompanyData, Sector
from generation.anonymizer import Anonymizer
from generation.llm_generator import ContentGenerator, TeaserContent
from ppt.ppt_generator import PresentationGenerator


def setup_argument_parser() -> argparse.ArgumentParser:
    """Set up command line argument parser."""
    parser = argparse.ArgumentParser(
        description="Kelp Automated Deal Flow - Investment Teaser Generator",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
  # Generate teaser for a specific company
  python main.py --company technology-ksolves

  # Generate teasers for all companies
  python main.py --all

  # List available companies
  python main.py --list

  # Debug mode with verbose output
  python main.py --company technology-ksolves --debug
        """
    )

    parser.add_argument(
        "--company", "-c",
        type=str,
        help="Company folder name (e.g., technology-ksolves)"
    )

    parser.add_argument(
        "--all", "-a",
        action="store_true",
        help="Process all companies"
    )

    parser.add_argument(
        "--list", "-l",
        action="store_true",
        help="List available companies"
    )

    parser.add_argument(
        "--output", "-o",
        type=str,
        help="Output directory (default: ./output)"
    )

    parser.add_argument(
        "--debug", "-d",
        action="store_true",
        help="Enable debug mode with verbose logging"
    )

    parser.add_argument(
        "--no-citations",
        action="store_true",
        help="Skip citation document generation"
    )

    parser.add_argument(
        "--test-parse",
        action="store_true",
        help="Only test data parsing (no PPT generation)"
    )

    return parser


def list_companies(data_loader: DataLoader) -> None:
    """List all available companies."""
    logger = get_logger("main")

    print("\n" + "="*60)
    print(" Available Companies")
    print("="*60 + "\n")

    data_dir = Path(data_loader.data_dir)
    company_folders = sorted([
        f.name for f in data_dir.iterdir()
        if f.is_dir() and not f.name.startswith('.')
    ])

    for folder in company_folders:
        # Try to extract company name
        onepager_files = list((data_dir / folder).glob("*-OnePager.md"))
        if onepager_files:
            name = onepager_files[0].stem.replace("-OnePager", "")
            print(f"  {folder:30} -> {name}")
        else:
            print(f"  {folder:30} -> (no OnePager found)")

    print(f"\nTotal: {len(company_folders)} companies")
    print("\nUsage: python main.py --company <folder_name>")


def process_company(
    company_data: CompanyData,
    config,
    test_parse_only: bool = False
) -> bool:
    """
    Process a single company and generate outputs.

    Args:
        company_data: Parsed company data
        config: Pipeline configuration
        test_parse_only: If True, only show parsed data

    Returns:
        True if successful
    """
    logger = get_logger("main")

    # Assign a project codename
    import hashlib
    name_hash = int(hashlib.md5(company_data.name.encode()).hexdigest(), 16)
    codename = PROJECT_CODENAMES[name_hash % len(PROJECT_CODENAMES)]

    logger.info(f"Processing: {company_data.name} -> Project {codename}")

    if test_parse_only:
        # Just display the parsed data
        display_parsed_data(company_data, codename)
        return True

    # Get sector template
    sector_template = get_sector_template(company_data.sector.value)
    logger.info(f"Using template for sector: {company_data.sector.value}")

    # Create anonymizer
    anonymizer = Anonymizer(company_data.name, codename)

    # Generate content using LLM
    logger.info("Generating content with LLM...")
    with LogTimer("Content generation", "main"):
        content_generator = ContentGenerator(config.llm)
        teaser_content = content_generator.generate_teaser_content(company_data, anonymizer)

        # Log usage stats
        usage = content_generator.get_usage_stats()
        logger.info(f"LLM Usage - Tokens: {usage['total_tokens']}, Cost: INR {usage['estimated_cost_inr']}")

    # Generate PPT
    logger.info("Generating PowerPoint presentation...")
    with LogTimer("PPT generation", "main"):
        ppt_generator = PresentationGenerator(config.branding)

        # Create output filename
        output_filename = f"Project_{codename}_{company_data.sector.value}.pptx"
        output_path = config.output_dir / output_filename

        ppt_path = ppt_generator.create_presentation(teaser_content, str(output_path))
        logger.info(f"PPT saved: {ppt_path}")

    # Generate citations document
    if config.generate_citations:
        logger.info("Generating citations document...")
        citations_path = generate_citations(company_data, teaser_content, config, codename)
        logger.info(f"Citations saved: {citations_path}")

    logger.info(f"Successfully processed: {company_data.name}")
    return True


def generate_citations(
    company_data: CompanyData,
    teaser_content: TeaserContent,
    config,
    codename: str
) -> str:
    """
    Generate a citations document for the teaser in Word format.

    Args:
        company_data: Original company data
        teaser_content: Generated teaser content
        config: Pipeline configuration
        codename: Project codename

    Returns:
        Path to citations document
    """
    from datetime import datetime

    try:
        from docx import Document
        from docx.shared import Pt, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        DOCX_AVAILABLE = True
    except ImportError:
        DOCX_AVAILABLE = False

    if DOCX_AVAILABLE:
        # Generate Word document
        citations_filename = f"Project_{codename}_Citations.docx"
        citations_path = config.output_dir / citations_filename

        doc = Document()

        # Title
        title = doc.add_heading(f'Citation Document - Project {codename}', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Generated date
        doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        doc.add_paragraph()

        # Data Sources section
        doc.add_heading('Data Sources', level=1)
        for source_name, source_path in company_data.sources.items():
            doc.add_paragraph(f"{source_name}: {source_path}", style='List Bullet')

        # Slide 1 citations
        doc.add_heading('Slide 1: Business Overview', level=1)
        citations_slide1 = [
            ("Business Description", "OnePager.md - Business Description section"),
            ("Founded", "OnePager.md - Details section"),
            ("Employees", f"OnePager.md - People section ({company_data.employees})"),
            ("Facilities", f"OnePager.md - Facilities section ({len(company_data.facilities)} locations)"),
            ("Products/Services", "OnePager.md - Product & Services section"),
            ("Certifications", "OnePager.md - Awards and Certifications section"),
        ]
        for label, source in citations_slide1:
            doc.add_paragraph(f"{label}: {source}", style='List Bullet')

        # Slide 2 citations
        doc.add_heading('Slide 2: Financial Metrics', level=1)
        latest = company_data.get_latest_financials()
        if latest:
            citations_slide2 = [
                (f"Revenue ({latest.year})", "OnePager.md - Financials Status - Income Statement"),
                (f"EBITDA ({latest.year})", "OnePager.md - Financials Status - Income Statement (Operating EBITDA)"),
                (f"PAT ({latest.year})", "OnePager.md - Financials Status - Income Statement"),
                ("PAT Margin", "OnePager.md - Financials Status - Ratios"),
                ("RoCE", "OnePager.md - Financials Status - Ratios"),
                ("ROE", "OnePager.md - Financials Status - Ratios"),
            ]
            for label, source in citations_slide2:
                doc.add_paragraph(f"{label}: {source}", style='List Bullet')

        # Slide 3 citations
        doc.add_heading('Slide 3: Investment Highlights', level=1)
        citations_slide3 = [
            ("Strengths", "OnePager.md - SWOT Analysis - Strengths section"),
            ("Opportunities", "OnePager.md - SWOT Analysis - Opportunities section"),
            ("Future Plans", "OnePager.md - Future Plan section"),
            ("Market Data", "OnePager.md - Market Size section"),
        ]
        for label, source in citations_slide3:
            doc.add_paragraph(f"{label}: {source}", style='List Bullet')

        # Market data details
        if company_data.market_data:
            doc.add_heading('Market Size Data Sources', level=1)
            for md in company_data.market_data[:5]:
                doc.add_paragraph(f"{md.market} ({md.region})", style='List Bullet')
                doc.add_paragraph(f"Source: {md.source}", style='List Bullet 2')
                doc.add_paragraph(f"Size: {md.current_size}", style='List Bullet 2')
                doc.add_paragraph(f"Growth Rate: {md.growth_rate}%", style='List Bullet 2')

        # Footer
        doc.add_paragraph()
        footer = doc.add_paragraph("This document was auto-generated by the Kelp Automated Deal Flow system.")
        footer.italic = True

        doc.save(str(citations_path))
        return str(citations_path)

    else:
        # Fallback to Markdown
        citations_filename = f"Project_{codename}_Citations.md"
        citations_path = config.output_dir / citations_filename

        with open(citations_path, 'w', encoding='utf-8') as f:
            f.write(f"# Citation Document - Project {codename}\n\n")
            f.write(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n\n")
            f.write("---\n\n")

            f.write("## Data Sources\n\n")
            for source_name, source_path in company_data.sources.items():
                f.write(f"- **{source_name}**: `{source_path}`\n")

            f.write("\n---\n\n")

            f.write("## Slide 1: Business Overview\n\n")
            f.write(f"- **Business Description**: OnePager.md - Business Description section\n")
            f.write(f"- **Founded**: OnePager.md - Details section\n")
            f.write(f"- **Employees**: OnePager.md - People section ({company_data.employees})\n")
            f.write(f"- **Facilities**: OnePager.md - Facilities section ({len(company_data.facilities)} locations)\n")
            f.write(f"- **Products/Services**: OnePager.md - Product & Services section\n")
            f.write(f"- **Certifications**: OnePager.md - Awards and Certifications section\n")

            f.write("\n---\n\n")

            f.write("## Slide 2: Financial Metrics\n\n")
            latest = company_data.get_latest_financials()
            if latest:
                f.write(f"- **Revenue ({latest.year})**: OnePager.md - Financials Status - Income Statement\n")
                f.write(f"- **EBITDA ({latest.year})**: OnePager.md - Financials Status - Income Statement\n")
                f.write(f"- **PAT ({latest.year})**: OnePager.md - Financials Status - Income Statement\n")
                f.write(f"- **PAT Margin**: OnePager.md - Financials Status - Ratios\n")
                f.write(f"- **RoCE**: OnePager.md - Financials Status - Ratios\n")
                f.write(f"- **ROE**: OnePager.md - Financials Status - Ratios\n")

            f.write("\n---\n\n")

            f.write("## Slide 3: Investment Highlights\n\n")
            f.write(f"- **Strengths**: OnePager.md - SWOT Analysis - Strengths section\n")
            f.write(f"- **Opportunities**: OnePager.md - SWOT Analysis - Opportunities section\n")
            f.write(f"- **Future Plans**: OnePager.md - Future Plan section\n")
            f.write(f"- **Market Data**: OnePager.md - Market Size section\n")

            f.write("\n---\n\n")

            if company_data.market_data:
                f.write("## Market Size Data Sources\n\n")
                for md in company_data.market_data[:5]:
                    f.write(f"- **{md.market}** ({md.region})\n")
                    f.write(f"  - Source: {md.source}\n")
                    f.write(f"  - Size: {md.current_size}\n")
                    f.write(f"  - Growth Rate: {md.growth_rate}%\n\n")

            f.write("\n---\n\n")
            f.write("*This document was auto-generated by the Kelp Automated Deal Flow system.*\n")

        return str(citations_path)


def display_parsed_data(company: CompanyData, codename: str) -> None:
    """Display parsed company data for verification."""
    # Set UTF-8 encoding for console output on Windows
    import sys
    import io
    if sys.platform == 'win32':
        sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

    print("\n" + "="*70)
    print(f" PARSED DATA: {company.name}")
    print(f" Project Codename: {codename}")
    print("="*70)

    print(f"\n[BASIC INFO]")
    print(f"  Sector:          {company.sector.value}")
    print(f"  Website:         {company.website}")
    print(f"  Founded:         {company.founded or 'N/A'}")
    print(f"  Headquarters:    {company.headquarters[:50] + '...' if len(company.headquarters) > 50 else company.headquarters}")
    print(f"  Domain:          {company.domain}")
    print(f"  Segment:         {company.segment}")
    print(f"  Customer Base:   {company.customer_base}")
    print(f"  Ownership:       {company.ownership_type}")

    print(f"\n[BUSINESS DESCRIPTION]")
    desc = company.business_description
    print(f"  {desc[:200]}..." if len(desc) > 200 else f"  {desc}")

    print(f"\n[PRODUCTS & SERVICES] ({len(company.products_services)} items)")
    for i, product in enumerate(company.products_services[:5]):
        print(f"  {i+1}. {product}")
    if len(company.products_services) > 5:
        print(f"  ... and {len(company.products_services) - 5} more")

    print(f"\n[INDUSTRIES SERVED]")
    print(f"  {', '.join(company.industries_served[:5])}")

    print(f"\n[KEY NUMBERS]")
    print(f"  Employees:       {company.employees or 'N/A'}")
    print(f"  Facilities:      {len(company.facilities)}")
    print(f"  Certifications:  {len(company.certifications)}")
    print(f"  Clients:         {len(company.clients)}")
    print(f"  Partners:        {len(company.partners)}")

    print(f"\n[GLOBAL PRESENCE]")
    print(f"  {', '.join(company.global_presence) if company.global_presence else 'N/A'}")

    print(f"\n[FACILITIES]")
    for facility in company.facilities[:5]:
        print(f"  - {facility.facility_type}: {facility.location}")

    print(f"\n[CERTIFICATIONS]")
    for cert in company.certifications[:5]:
        print(f"  - {cert}")

    print(f"\n[SWOT ANALYSIS]")
    print(f"  Strengths:      {len(company.swot.strengths)}")
    print(f"  Weaknesses:     {len(company.swot.weaknesses)}")
    print(f"  Opportunities:  {len(company.swot.opportunities)}")
    print(f"  Threats:        {len(company.swot.threats)}")

    if company.swot.strengths:
        print(f"\n  Top Strength:")
        print(f"    {company.swot.strengths[0][:100]}...")

    print(f"\n[FINANCIALS]")
    latest = company.get_latest_financials()
    if latest:
        print(f"  Latest Year:     {latest.year}")
        print(f"  Revenue:         {latest.revenue or 'N/A'}")
        print(f"  EBITDA:          {latest.ebitda or 'N/A'}")
        print(f"  PAT:             {latest.pat or 'N/A'}")
        print(f"  PAT Margin:      {latest.pat_margin or 'N/A'}%")
        print(f"  RoCE:            {latest.roce or 'N/A'}%")
        print(f"  ROE:             {latest.roe or 'N/A'}%")

        # Calculate CAGR
        revenue_cagr = company.financials.calculate_cagr('revenue', 5)
        if revenue_cagr:
            print(f"  Revenue CAGR (5Y): {revenue_cagr:.1f}%")
    else:
        print("  No financial data available")

    print(f"\n[MILESTONES] ({len(company.milestones)} total)")
    for milestone in company.milestones[:5]:
        print(f"  {milestone.date}: {milestone.description[:60]}...")

    print(f"\n[FUTURE PLANS]")
    for plan in company.future_plans[:3]:
        print(f"  - {plan[:80]}...")

    print("\n" + "="*70)


def main():
    """Main entry point."""
    # Parse arguments
    parser = setup_argument_parser()
    args = parser.parse_args()

    # Setup logging
    log_level = 10 if args.debug else 20  # DEBUG or INFO
    setup_logger("main", level=log_level)
    setup_logger("data_loader", level=log_level)
    setup_logger("markdown_parser", level=log_level)

    logger = get_logger("main")

    logger.info("="*60)
    logger.info(" Kelp Automated Deal Flow")
    logger.info(" AI-ML GC 2025-26")
    logger.info("="*60)

    # Get configuration
    config = get_config()

    # Override output directory if specified
    if args.output:
        config.output_dir = Path(args.output)
        config.output_dir.mkdir(parents=True, exist_ok=True)

    # Initialize data loader
    data_loader = DataLoader(str(config.data_dir))

    # Handle list command
    if args.list:
        list_companies(data_loader)
        return 0

    # Check if company or all flag is provided
    if not args.company and not args.all:
        parser.print_help()
        print("\nError: Please specify --company or --all")
        return 1

    # Process companies
    if args.all:
        logger.info("Processing all companies...")
        companies = data_loader.load_all_companies()

        success_count = 0
        for name, company_data in companies.items():
            try:
                if process_company(company_data, config, args.test_parse):
                    success_count += 1
            except Exception as e:
                logger.error(f"Failed to process {name}: {e}")

        logger.info(f"Processed {success_count}/{len(companies)} companies successfully")

    else:
        # Process single company
        logger.info(f"Processing company: {args.company}")

        company_data = data_loader.load_company(args.company)
        if company_data is None:
            logger.error(f"Could not load company: {args.company}")
            return 1

        try:
            process_company(company_data, config, args.test_parse)
        except Exception as e:
            logger.error(f"Failed to process {args.company}: {e}")
            if args.debug:
                raise
            return 1

    logger.info("Done!")
    return 0


if __name__ == "__main__":
    sys.exit(main())
